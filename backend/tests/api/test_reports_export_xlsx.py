"""Tests for the server-side XLSX export routes (Phase B / B2+B3).

Exercises the per-report ``export.xlsx`` endpoints through the HTTP layer
(same ``reports:read`` + RLS scoping as the JSON siblings), re-opens the
streamed bytes with openpyxl, and asserts the structure the report-export
service is supposed to produce: bold frozen headers, performance fills equal
to the ``colors.ts`` hex, freeze panes, and values matching the payload.
"""

from __future__ import annotations

import io

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.schemas.reports import (
    QuestionResponseAnalysisPayload,
    QuestionSummaryMatrixPayload,
    StandardSummaryPayload,
)
from app.services import report_export_service as X

_XLSX_MEDIA = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
_CHAPTER9 = "8359960427"


def _load(content: bytes):
    return load_workbook(io.BytesIO(content))


@pytest.mark.anyio
async def test_qra_export_xlsx_structure_and_fills(
    admin_client: AsyncClient,
) -> None:
    # Fetch the JSON payload so we can compare cell values.
    j = await admin_client.get(
        f"/api/v1/reports/question-response-analysis/{_CHAPTER9}"
    )
    if j.status_code == 404:
        pytest.skip("Chapter 9 fixture item missing")
    payload = QuestionResponseAnalysisPayload.model_validate(j.json())

    r = await admin_client.get(
        f"/api/v1/reports/question-response-analysis/{_CHAPTER9}/export.xlsx"
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(_XLSX_MEDIA)
    cd = r.headers["content-disposition"]
    assert cd.startswith("attachment; filename=")
    # No CR/LF/quote injection survived into the header.
    assert "\r" not in cd and "\n" not in cd

    wb = _load(r.content)
    ws = wb.active
    assert ws.title == "QRA"
    assert [ws.cell(1, c).value for c in range(1, 9)] == X._QRA_HEADERS
    assert ws.cell(1, 1).font.bold is True
    assert ws.cell(1, 1).fill.fgColor.rgb == X.HEADER_BAR_BG
    assert ws.freeze_panes == "A2"
    assert ws.max_row == len(payload.questions_overall) + 1

    q0 = payload.questions_overall[0]
    cell = ws.cell(2, 5)  # Grade Average %
    assert cell.value == pytest.approx(q0.grade_average)
    assert cell.number_format == "0.0%"
    expected = (
        X.PERF_PINK
        if q0.grade_average < 0.7
        else X.PERF_YELLOW
        if q0.grade_average < 0.8
        else X.PERF_GREEN
    )
    assert cell.fill.fgColor.rgb == expected


@pytest.mark.anyio
async def test_qsr_export_xlsx_matrix_bands_and_fills(
    admin_client: AsyncClient,
) -> None:
    j = await admin_client.get(
        f"/api/v1/reports/question-summary-paginated/{_CHAPTER9}"
    )
    if j.status_code == 404:
        pytest.skip("Chapter 9 fixture item missing")
    payload = QuestionSummaryMatrixPayload.model_validate(j.json())

    r = await admin_client.get(
        f"/api/v1/reports/question-summary-paginated/{_CHAPTER9}/export.xlsx"
    )
    assert r.status_code == 200, r.text
    wb = _load(r.content)
    ws = wb.active
    assert ws.title == "Question Summary"

    nq = len(payload.questions)
    # Band header row 1, navy.
    assert ws.cell(1, 1).value == "Standards"
    assert ws.cell(1, 1).fill.fgColor.rgb == X.PBIX_ACCENT_NAVY
    # Column header row 2, leaf question_no order.
    assert ws.cell(2, 1).value == "Classroom Instructors"
    assert ws.cell(2, 3).value == "Score %"
    leaf = [ws.cell(2, 4 + i).value for i in range(nq)]
    assert leaf == [q.question_no for q in payload.questions]
    assert ws.cell(2, 4 + nq).value == "Possible Points"
    assert ws.cell(2, 4 + nq + 1).value == "# Correct Answers"
    # Freeze at body start (D3).
    assert ws.freeze_panes == ws.cell(3, 4).coordinate

    # First student's first attempted binary cell carries a QSR fill.
    s0 = payload.teacher_groups[0].students[0]
    for i, q in enumerate(payload.questions):
        cv = s0.cells.get(q.question_id)
        if cv in (0, 1):
            cell = ws.cell(3, 4 + i)
            assert cell.value == cv
            assert cell.fill.fgColor.rgb == (X.QSR_GREEN if cv == 1 else X.QSR_PINK)
            break

    # Grand-total trio labels at the foot.
    labels = [ws.cell(ws.max_row - 2 + k, 1).value for k in range(3)]
    assert labels == ["Possible Points", "# Correct Answers", "Score %"]


@pytest.mark.anyio
async def test_standard_summary_export_xlsx(admin_client: AsyncClient) -> None:
    j = await admin_client.get("/api/v1/reports/standard-summary")
    assert j.status_code == 200, j.text
    payload = StandardSummaryPayload.model_validate(j.json())

    r = await admin_client.get("/api/v1/reports/standard-summary/export.xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(_XLSX_MEDIA)
    wb = _load(r.content)
    ws = wb.active
    assert ws.freeze_panes == "A2"
    assert ws.cell(1, 1).value == "Standard"
    assert ws.max_row == len(payload.standards) + 1
    if payload.standards:
        s0 = payload.standards[0]
        assert ws.cell(2, 1).value == s0.schoology_standard
        cell = ws.cell(2, 7)  # Grade Average %
        assert cell.value == pytest.approx(s0.grade_average)


@pytest.mark.anyio
async def test_export_xlsx_requires_reports_read(
    user_client: AsyncClient,
) -> None:
    r = await user_client.get(
        f"/api/v1/reports/question-response-analysis/{_CHAPTER9}/export.xlsx"
    )
    assert r.status_code == 403


@pytest.mark.anyio
async def test_export_xlsx_unknown_item_404(admin_client: AsyncClient) -> None:
    r = await admin_client.get(
        "/api/v1/reports/question-response-analysis/0000000000/export.xlsx"
    )
    assert r.status_code == 404


def test_perf_and_qsr_fill_hex_match_colors_ts() -> None:
    """The XLSX fill hexes MUST equal the colors.ts literals (trailing 6 chars).

    Guards against drift between the exporter palette and the single source of
    truth in ``frontend/src/lib/reports/colors.ts``.
    """
    import pathlib
    import re

    colors_ts = pathlib.Path(__file__).resolve().parents[3] / (
        "frontend/src/lib/reports/colors.ts"
    )
    text = colors_ts.read_text()

    def literal(name: str) -> str:
        m = re.search(rf"export const {name} = '#([0-9A-Fa-f]{{6}})'", text)
        assert m, f"{name} not found in colors.ts"
        return m.group(1).upper()

    assert X.PERF_PINK[2:] == literal("PERF_PINK")
    assert X.PERF_YELLOW[2:] == literal("PERF_YELLOW")
    assert X.PERF_GREEN[2:] == literal("PERF_GREEN")
    assert X.QSR_PINK[2:] == literal("QSR_PINK")
    assert X.QSR_YELLOW[2:] == literal("QSR_YELLOW")
    assert X.QSR_GREEN[2:] == literal("QSR_GREEN")
    assert X.HEADER_BAR_BG[2:] == literal("HEADER_BAR_BG")
    assert X.LAYOUT_BORDER[2:] == literal("LAYOUT_BORDER")
    assert X.PBIX_ACCENT_NAVY[2:] == literal("PBIX_ACCENT_NAVY")


def test_sanitize_xlsx_filename_strips_header_injection() -> None:
    bad = 'Chapter\r\n9 "Test"/<x>'
    out = X.sanitize_xlsx_filename("qra", bad)
    assert out.endswith(".xlsx")
    assert "\r" not in out and "\n" not in out and '"' not in out
    assert "/" not in out
    assert out.startswith("qra-")
    # Empty / None item name falls back to the report kind.
    assert X.sanitize_xlsx_filename("qsr", None) == "qsr-qsr.xlsx"
